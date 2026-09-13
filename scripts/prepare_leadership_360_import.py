#!/usr/bin/env python3
"""Convert the approved Leadership 360 XLSX export into a private import file.

The output contains raw response material. Write it only beneath /work/ or
another restricted location; those paths are intentionally ignored by Git.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SCORE_VALUES = {
    "Highly ineffectively": 1,
    "Somewhat ineffectively": 2,
    "Neutral": 3,
    "Somewhat effectively": 4,
    "Highly effectively": 5,
}
QUESTION_IDS = [f"Q{number:02d}" for number in range(3, 19)]
LEADER_COLUMN = "Which SLG member are you assessing?"
RELATIONSHIP_COLUMN = "What is your relationship with this SLG member?"
RESPONSE_ID_COLUMN = "Response ID"
OPEN_FIELDS = {
    "What are this SLG members greatest strengths?": "strengths",
    "What are this SLG members most important development areas?": "development",
    "Anything else you would like to share about this SLG member?": "otherFeedback",
}


def manager_id(name: str) -> str:
    digest = hashlib.sha256(name.casefold().encode("utf-8")).hexdigest()[:16]
    return f"manager-{digest}"


def report_period(values: list[object]) -> str:
    dates = []
    for value in values:
        if isinstance(value, datetime):
            dates.append(value)
        elif isinstance(value, str):
            dates.append(datetime.strptime(value, "%d %B %Y %H:%M:%S"))
    if not dates:
        return "Leadership 360"
    start, end = min(dates), max(dates)
    if start.year == end.year and start.month == end.month:
        return start.strftime("%B %Y")
    if start.year == end.year:
        return f"{start.strftime('%B')}–{end.strftime('%B %Y')}"
    return f"{start.strftime('%B %Y')}–{end.strftime('%B %Y')}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The workbook has no rows.")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    positions = {header: index for index, header in enumerate(headers)}
    for required in (RESPONSE_ID_COLUMN, LEADER_COLUMN, RELATIONSHIP_COLUMN):
        if required not in positions:
            raise ValueError(f"Missing required column: {required}")

    score_headers = [
        header for header in headers
        if header.startswith("How effectively does this SLG member") and not header.endswith(": Comments")
    ]
    if len(score_headers) != len(QUESTION_IDS):
        raise ValueError(f"Expected {len(QUESTION_IDS)} scored questions, found {len(score_headers)}.")

    managers: OrderedDict[str, dict[str, object]] = OrderedDict()
    responses = []
    submitted = []
    date_index = positions.get("Date Submitted")

    for row in rows[1:]:
        leader = str(row[positions[LEADER_COLUMN]] or "").strip()
        relationship = str(row[positions[RELATIONSHIP_COLUMN]] or "").strip()
        response_id = str(row[positions[RESPONSE_ID_COLUMN]] or "").strip()
        if not leader or not relationship or not response_id:
            raise ValueError("Every response needs an ID, assessed leader and relationship.")
        identifier = manager_id(leader)
        managers.setdefault(identifier, {
            "id": identifier,
            "name": leader,
            "jobTitle": "Senior Leadership Group",
        })
        scores = {}
        question_comments = {}
        for question_id, header in zip(QUESTION_IDS, score_headers, strict=True):
            raw_score = row[positions[header]]
            if raw_score in (None, ""):
                scores[question_id] = None
            elif raw_score in SCORE_VALUES:
                scores[question_id] = SCORE_VALUES[raw_score]
            else:
                raise ValueError(f"Unexpected score '{raw_score}' in response {response_id}.")
            comment = row[positions[f"{header}: Comments"]]
            question_comments[question_id] = str(comment).strip() if comment else None

        response = {
            "managerId": identifier,
            "responseId": response_id,
            "relationship": relationship,
            "scores": scores,
            "questionComments": question_comments,
        }
        for source_header, target_key in OPEN_FIELDS.items():
            value = row[positions[source_header]]
            response[target_key] = str(value).strip() if value else None
        responses.append(response)
        if date_index is not None and row[date_index]:
            submitted.append(row[date_index])

    period = report_period(submitted)
    for manager in managers.values():
        manager["reportPeriod"] = period

    checksum = hashlib.file_digest(args.source.open("rb"), "sha256").hexdigest()
    payload = {
        "schemaVersion": 1,
        "sourceFilename": args.source.name,
        "sourceChecksum": checksum,
        "reportPeriod": period,
        "managers": list(managers.values()),
        "responses": responses,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(f"Prepared {len(managers)} managers and {len(responses)} responses.")


if __name__ == "__main__":
    main()
